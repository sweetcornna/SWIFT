from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "project-plan.md"
OUTPUT = ROOT / "docs" / "project-plan.docx"


def set_run_font(run, name: str = "Microsoft YaHei", size: int | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    r_fonts = run._element.rPr.rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top: int = 120, start: int = 120, bottom: int = 120, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "1F4D78", 14, 8),
        ("Heading 2", 13, "2E74B5", 10, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("低空闪翼项目计划书")
    set_run_font(run, size=9)
    run.font.color.rgb = RGBColor.from_string("666666")


def add_paragraph_with_inline_code(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    tokens = re.split(r"(`[^`]+`)", text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run(token[1:-1] if token.startswith("`") and token.endswith("`") else token)
        set_run_font(run)
        if token.startswith("`") and token.endswith("`"):
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.color.rgb = RGBColor.from_string("1F4D78")
    return paragraph


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value.strip())
            set_run_font(run, size=9)
            if row_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("0B2545")
                set_cell_shading(cell, "E8EEF5")
    doc.add_paragraph()


def parse_markdown_tables(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        if not all(set(cell) <= {"-", ":", " "} for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def build_docx() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)

    lines = text.splitlines()
    index = 0
    first_title = True
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue

        if line.startswith("|"):
            rows, index = parse_markdown_tables(lines, index)
            add_table(doc, rows)
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(12)
            run = paragraph.add_run(title)
            set_run_font(run, size=20)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("0B2545")
            if first_title:
                first_title = False
            index += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            index += 1
            continue

        if line.startswith("- "):
            paragraph = add_paragraph_with_inline_code(doc, line[2:].strip(), style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.25)
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            paragraph = add_paragraph_with_inline_code(doc, numbered.group(2), style="List Number")
            paragraph.paragraph_format.left_indent = Inches(0.25)
            index += 1
            continue

        paragraph = add_paragraph_with_inline_code(doc, line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        index += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT)
